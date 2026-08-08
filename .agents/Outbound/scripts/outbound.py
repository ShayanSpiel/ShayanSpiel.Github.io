#!/usr/bin/env python3
"""
SpielOS Outbound Email Automation
Reads the master outreach database and sends personalized emails via a
provider selected by env vars (resend | sendgrid | mailgun | smtp).

Usage:
  python3 outbound.py stats                       # Database stats
  python3 outbound.py dry-run [--lang en|fa] [--tier A] [--limit 5]
  python3 outbound.py send --limit 10 [--lang en] [--tier A] [--retry-failed]
  python3 outbound.py send --retry-failed                 # re-send failed[] entries
  python3 outbound.py metrics [--force] [--quiet] # Email Data: pull provider status (scheduled)
  python3 outbound.py review [--json]           # Goal check vs reply-rate goal + next action
  python3 outbound.py record-reply <email|lead_id> [--note "..."]  # log a reply from the inbox
  python3 outbound.py replies                     # list recorded replies

Config:
  All configuration lives in .agents/Outbound/.env (see .env.example):
  provider keys, sender identity, signature, list path, throttle, rotation,
  metrics cadence, and funnel goals.
  Templates live in templates.py (language -> list of A/B variants).
"""

import argparse
import json
import os
import sys
import time
from datetime import datetime, timezone

import analytics
import config
import providers
import strategy
from templates import TEMPLATES, SIGNATURE_HTML, SIGNATURE_TEXT

config.validate()

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip3 install openpyxl")
    sys.exit(1)

SHEET_NAME = config.SHEET_NAME

COL_MAP = {
    "lead_id": 0,
    "send_recommendation": 1,
    "outreach_tier": 2,
    "company_contact_rank": 3,
    "contactability": 4,
    "market": 5,
    "company": 6,
    "company_domain": 7,
    "contact_name": 8,
    "title": 9,
    "email": 10,
    "email_status": 11,
    "person_linkedin": 12,
    "website": 13,
    "segment": 14,
    "country": 15,
    "employees": 16,
    "annual_revenue": 17,
    "technologies": 18,
    "need_buying_signals": 19,
    "icp_confidence": 20,
    "qualification_rationale": 21,
    "pain_hypothesis": 22,
    "recommended_pilot": 23,
    "personalization_hook": 24,
    "suggested_cta": 25,
    "language": 26,
    "source": 27,
    "source_url": 28,
    "agent_instructions": 29,
    "sequence_status": 30,
    "last_checked": 31,
    "notes": 32,
    "apollo_contact_id": 33,
    "apollo_account_id": 34,
}

LANG_ALIASES = {"en": "English", "fa": "Persian", "english": "English", "persian": "Persian"}


# ── Helpers ────────────────────────────────────────────────────────────────────

def load_sent_log() -> dict:
    if config.SENT_LOG_PATH.exists():
        with open(config.SENT_LOG_PATH) as f:
            return json.load(f)
    return {"sent": [], "failed": []}


def save_sent_log(log: dict):
    """Atomic write (tmp + rename) — a stale batch child must never leave a
    torn file behind for the next reader."""
    tmp = config.SENT_LOG_PATH.with_suffix(".json.tmp")
    with open(tmp, "w") as f:
        json.dump(log, f, indent=2, default=str)
    os.replace(tmp, config.SENT_LOG_PATH)


def already_sent(lead_id: str, log: dict) -> bool:
    return any(s.get("lead_id") == lead_id for s in log.get("sent", []))


def read_contacts(lang_filter=None, tier_filter=None):
    wb = openpyxl.load_workbook(config.DATABASE_PATH, read_only=True, data_only=True)
    ws = wb[SHEET_NAME]
    contacts = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if not row or not row[COL_MAP["email"]]:
            continue
        contact = {k: row[v] for k, v in COL_MAP.items()}
        contact["_row"] = i + 2

        contact["language"] = str(contact.get("language") or "English").strip()
        contact["email"] = str(contact["email"]).strip().lower()
        contact["company"] = str(contact.get("company") or "").strip()
        contact["contact_name"] = str(contact.get("contact_name") or "").strip()
        contact["title"] = str(contact.get("title") or "").strip()
        contact["domain"] = str(contact.get("company_domain") or "").strip()
        contact["website"] = str(contact.get("website") or "").strip()
        contact["country"] = str(contact.get("country") or "").strip()
        contact["segment"] = str(contact.get("segment") or "").strip()
        contact["personalization_hook"] = str(contact.get("personalization_hook") or "").strip()
        contact["suggested_cta"] = str(contact.get("suggested_cta") or "").strip()
        contact["lead_id"] = str(contact.get("lead_id") or "").strip()
        contact["sequence_status"] = str(contact.get("sequence_status") or "").strip()
        contact["send_recommendation"] = str(contact.get("send_recommendation") or "").strip()
        contact["outreach_tier"] = str(contact.get("outreach_tier") or "").strip()
        contact["email_status"] = str(contact.get("email_status") or "").strip()

        if lang_filter and contact["language"].lower() != lang_filter.lower():
            continue
        if tier_filter and contact["outreach_tier"].upper() != tier_filter.upper():
            continue

        contacts.append(contact)

    wb.close()
    return contacts


def get_first_name(contact: dict) -> str:
    name = contact["contact_name"]
    if not name:
        return ""
    return name.strip().split()[0]


def pick_variant(lang: str, index: int):
    """Rotate A/B variants: every VARIANT_ROTATE emails, rotate to the next variant."""
    variants = TEMPLATES.get(lang, TEMPLATES["English"])
    return variants[(index // config.VARIANT_ROTATE) % len(variants)]


def render_template(template_str: str, contact: dict) -> str:
    first_name = get_first_name(contact)
    if not first_name:
        first_name = "there"

    return template_str.format(
        contact_name=contact["contact_name"],
        first_name=first_name,
        company=contact["company"],
        title=contact["title"],
        domain=contact["domain"],
        personalization_hook=contact["personalization_hook"],
        suggested_cta=contact["suggested_cta"],
        website=contact["website"],
        country=contact["country"],
        segment=contact["segment"],
        SIGNATURE_HTML=SIGNATURE_HTML,
        SIGNATURE_TEXT=SIGNATURE_TEXT,
    )


def templates_ready() -> bool:
    return all(
        "TODO" not in v["subject"] and "TODO" not in v["body_html"]
        for variants in TEMPLATES.values()
        for v in variants
    )


# ── Commands ───────────────────────────────────────────────────────────────────

def cmd_stats():
    contacts = read_contacts()
    langs = {}
    tiers = {}
    statuses = {}
    recs = {}
    for c in contacts:
        langs[c["language"]] = langs.get(c["language"], 0) + 1
        tiers[c["outreach_tier"]] = tiers.get(c["outreach_tier"], 0) + 1
        statuses[c["sequence_status"]] = statuses.get(c["sequence_status"], 0) + 1
        recs[c["send_recommendation"]] = recs.get(c["send_recommendation"], 0) + 1

    log = load_sent_log()
    sent_count = len(log.get("sent", []))

    print(f"\n{'='*55}")
    print(f"  DATABASE STATS — Master Outreach")
    print(f"  Provider: {config.EMAIL_PROVIDER}")
    print(f"{'='*55}")
    print(f"  Total contacts:       {len(contacts)}")
    print(f"  Already sent:         {sent_count}")
    print(f"  Remaining:            {len(contacts) - sent_count}")
    print(f"")
    print(f"  By Language:")
    for k, v in sorted(langs.items()):
        print(f"    {k:<12} {v}")
    print(f"")
    print(f"  By Tier:")
    for k, v in sorted(tiers.items()):
        print(f"    {k:<12} {v}")
    print(f"")
    print(f"  By Sequence Status:")
    for k, v in sorted(statuses.items()):
        print(f"    {k:<12} {v}")
    print(f"")
    print(f"  By Send Recommendation:")
    for k, v in sorted(recs.items()):
        print(f"    {k:<12} {v}")
    print(f"{'='*55}\n")


def cmd_dry_run(lang_filter=None, tier_filter=None, limit=None):
    contacts = read_contacts(lang_filter=lang_filter, tier_filter=tier_filter)
    log = load_sent_log()

    unsent = [c for c in contacts if not already_sent(c["lead_id"], log)]
    if limit:
        unsent = unsent[:limit]

    print(f"\n{'='*60}")
    print(f"  DRY RUN — {len(unsent)} emails would be sent")
    if lang_filter:
        print(f"  Language filter: {lang_filter}")
    if tier_filter:
        print(f"  Tier filter: {tier_filter}")
    print(f"  Provider: {config.EMAIL_PROVIDER}")
    print(f"{'='*60}")

    if not templates_ready():
        print("\n  ⚠️  WARNING: Templates still contain 'TODO' placeholders.")
        print("  Edit templates.py before sending real emails.\n")

    for i, c in enumerate(unsent):
        lang = c["language"]
        tmpl = pick_variant(lang, i)
        subject = render_template(tmpl["subject"], c)
        body = render_template(tmpl["body_html"], c)

        print(f"\n  ── Email {i+1}/{len(unsent)} ──")
        print(f"  To:       {c['contact_name']} <{c['email']}>")
        print(f"  Company:  {c['company']} ({c['domain']})")
        print(f"  Tier:     {c['outreach_tier']} | Lang: {lang} | ICP: {c.get('icp_confidence', '?')}%")
        print(f"  Variant:  {tmpl['label']}")
        print(f"  Subject:  {subject}")
        print(f"  Body (first 200 chars):")
        print(f"    {body[:200]}...")
        print(f"  Status:   WOULD SEND (dry run)")

    print(f"\n{'='*60}")
    print(f"  End of dry run. {len(unsent)} emails previewed.")
    print(f"  To actually send: python3 outbound.py send --limit {len(unsent)}")
    if lang_filter:
        print(f"    Add: --lang {lang_filter}")
    print(f"{'='*60}\n")


def sent_today(log: dict, now=None) -> int:
    """Count sent entries whose timestamp is on the same UTC day as now."""
    now = now or datetime.now(timezone.utc)
    day = now.strftime("%Y-%m-%d")
    return sum(1 for s in log.get("sent", []) if str(s.get("timestamp", "")).startswith(day))


def daily_cap(now=None) -> tuple:
    """Deterministic daily send cap: warmup ramp by account age, hard-capped by
    the sum of the enabled providers' free-plan daily limits. Returns (cap, phase).

    Account age = days since the first logged send (sent_log.json).
    Phase ceilings are env-driven (WARMUP/RAMP/STEADY); the REAL ceiling is
    PROVIDER_DAILY_TOTAL (resend 100 + mailgun 100 + brevo 300 = 500/day once
    brevo is wired; sendgrid +100 when keyed). Halting is per-provider + the
    bounce/spam/delivery safety nets in analytics.py.
    """
    now = now or datetime.now(timezone.utc)
    log = load_sent_log()
    sent = log.get("sent", [])
    if not sent:
        cap = min(config.WARMUP_DAILY_CAP, config.PROVIDER_DAILY_TOTAL)
        return cap, f"warmup (<={cap}/day, no history)"
    first = _parse_ts(sent[0].get("timestamp"))
    if first is None:
        cap = min(config.WARMUP_DAILY_CAP, config.PROVIDER_DAILY_TOTAL)
        return cap, f"warmup (<={cap}/day)"
    age_days = (now - first).total_seconds() / 86400.0
    if age_days <= 14:
        cap = min(config.WARMUP_DAILY_CAP, config.PROVIDER_DAILY_TOTAL)
        phase = f"warmup day {age_days:.0f}/14 (<={cap}/day)"
    elif age_days <= 28:
        cap = min(config.RAMP_DAILY_CAP, config.PROVIDER_DAILY_TOTAL)
        phase = f"ramp day {age_days:.0f}/28 (<={cap}/day)"
    else:
        cap = min(config.STEADY_DAILY_CAP, config.PROVIDER_DAILY_TOTAL)
        phase = f"steady (<={cap}/day, provider hard caps)"
    return cap, phase


def _parse_ts(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None


def cmd_send(limit: int, lang_filter=None, tier_filter=None, retry_failed=False):
    contacts = read_contacts(lang_filter=lang_filter, tier_filter=tier_filter)
    log = load_sent_log()
    by_id = {c["lead_id"]: c for c in contacts}

    cap, phase = daily_cap()
    used_today = sent_today(log)
    remaining_today = cap - used_today
    if remaining_today <= 0:
        print(f"❌ Daily send cap reached ({used_today}/{cap} today, phase: {phase}).")
        print(f"   Next batch may run after midnight UTC, or raise WARMUP_DAILY_CAP after warmup.")
        return

    if retry_failed:
        failed_leads = [f["lead_id"] for f in log.get("failed", [])]
        to_send = [by_id[l] for l in failed_leads if l in by_id and not already_sent(l, log)]
        if not to_send:
            print("No failed sends to retry (all failed leads are already sent or not in the list).")
            return
        retry_ids = {c["lead_id"] for c in to_send}
        log["failed"] = [f for f in log.get("failed", []) if f.get("lead_id") not in retry_ids]
        save_sent_log(log)
        print(f"Retrying {len(to_send)} failed sends from sent_log.json...")
    else:
        unsent = [c for c in contacts if not already_sent(c["lead_id"], log)]
        batch = min(limit or config.BATCH_SIZE, config.BATCH_SIZE, remaining_today)
        to_send = unsent[:batch]
        if len(unsent) > batch:
            print(f"  (queue has {len(unsent)} unsent; sending batch of {batch} per cap {phase})")

    if not to_send:
        print("No contacts to send to (all filtered or already sent).")
        return

    if not templates_ready():
        print("\n❌ ABORT: Templates still contain 'TODO' placeholders.")
        print("Edit templates.py before sending real emails.\n")
        sys.exit(1)

    print(f"\n{'='*60}")
    print(f"  SENDING {len(to_send)} EMAILS via {config.EMAIL_PROVIDER}")
    print(f"  Rate: 1 email every {config.THROTTLE_SECONDS}s")
    print(f"  Daily cap: {used_today + len(to_send)}/{cap} after this batch ({phase})")
    print(f"  Variant rotation: every {config.VARIANT_ROTATE} emails")
    if config.REPLY_TO:
        print(f"  Reply-To: {config.REPLY_TO} (replies auto-detected by `metrics`)")
    print(f"  Est. time: ~{len(to_send) * config.THROTTLE_SECONDS:.0f}s ({len(to_send) * config.THROTTLE_SECONDS / 60:.1f} min)")
    print(f"{'='*60}\n")

    sent_count = 0
    fail_count = 0

    for i, c in enumerate(to_send):
        lang = c["language"]
        tmpl = pick_variant(lang, i)
        subject = render_template(tmpl["subject"], c)
        body_html = render_template(tmpl["body_html"], c)
        body_text = render_template(tmpl["body_text"], c)

        print(f"  [{i+1}/{len(to_send)}] Sending to {c['contact_name']} <{c['email']}> ({c['company']})...", end=" ", flush=True)

        result = providers.send_email(c["email"], subject, body_html, body_text, reply_to=config.REPLY_TO)

        if result.get("error"):
            print(f"❌ FAILED — {result.get('message', 'unknown')}")
            err_msg = str(result.get("message", ""))
            if "daily_quota_exceeded" in err_msg or "monthly_quota_exceeded" in err_msg or result.get("status") == 429:
                print("\n❌ Provider quota exceeded (429). Stopping the batch — do not send again "
                      "until the quota resets (daily at UTC midnight, monthly on plan renewal).")
                break
            log.setdefault("failed", []).append({
                "lead_id": c["lead_id"],
                "email": c["email"],
                "company": c["company"],
                "error": result.get("message", "unknown"),
                "timestamp": datetime.utcnow().isoformat(),
            })
            fail_count += 1
        else:
            print(f"✅ id={result.get('id', '?')[:12]}...")
            log.setdefault("sent", []).append({
                "lead_id": c["lead_id"],
                "email": c["email"],
                "company": c["company"],
                "contact_name": c["contact_name"],
                "variant": tmpl["label"],
                "subject": subject,
                "provider": config.EMAIL_PROVIDER,
                "provider_id": result.get("id"),
                "resend_id": result.get("id"),
                "timestamp": datetime.utcnow().isoformat(),
            })
            sent_count += 1

        if i < len(to_send) - 1:
            time.sleep(config.THROTTLE_SECONDS)

    save_sent_log(log)

    print(f"\n{'='*60}")
    print(f"  COMPLETE — Sent: {sent_count}, Failed: {fail_count}")
    print(f"  Log saved to: {config.SENT_LOG_PATH}")
    print(f"  Email Data: python3 outbound.py metrics --force")
    print(f"{'='*60}\n")


# ── Email Data commands ───────────────────────────────────────────────────────

def cmd_metrics(force=False, quiet=False):
    log = load_sent_log()
    if not log.get("sent"):
        print("No sent emails yet — nothing to measure. Run `send` first.")
        return

    if not analytics.cap_status_supported():
        print(f"\n  Provider '{config.EMAIL_PROVIDER}' does not report email status — "
              f"Email Data is unavailable for it.\n"
              f"  Sending still works; replies can be recorded manually "
              f"(`record-reply <email|lead_id>`).\n"
              f"  For full Email Data use resend or mailgun (see SKILL.md Part 4).\n")
        return

    metrics, ran = analytics.collect(log, force=force)
    if not ran:
        last = metrics.get("last_check") or "never"
        print(f"Metrics not due — last check {last}. Next pull in "
              f"{config.METRICS_INTERVAL_HOURS:g}h, or use --force.")
        return

    if quiet:
        t = analytics.aggregate(log, metrics)
        print(f"Email data refreshed ({t['sent']} checked): delivered {t['delivered_rate']*100:.0f}% · "
              f"opened {t['open_rate']*100:.0f}% · replied {t['reply_rate']*100:.0f}%")
    else:
        analytics.print_report(log, metrics)


def cmd_review(as_json=False):
    log = load_sent_log()
    metrics = analytics.load_metrics()
    if not log.get("sent"):
        print("No sent emails yet. Run `send` first, then `metrics`, then `review`.")
        return
    rep = strategy.review(log, metrics)
    if as_json:
        print(json.dumps(strategy.summary(rep), indent=2, default=str))
    else:
        strategy.print_review(rep)


def cmd_record_reply(identifier, note=None):
    log = load_sent_log()
    s = next(
        (x for x in log.get("sent", []) if x.get("lead_id") == identifier or x.get("email") == identifier),
        None,
    )
    if not s:
        print(f"ERROR: no sent email found for '{identifier}' (use lead_id or email).")
        sys.exit(1)

    metrics = analytics.load_metrics()
    for r in metrics.get("replies", []):
        if r["lead_id"] == s["lead_id"]:
            at = r.get("recorded_at") or r.get("received_at") or "?"
            print(f"ERROR: reply already recorded for {s['email']} on {at}.")
            sys.exit(1)

    metrics.setdefault("replies", []).append({
        "received_id": None,
        "lead_id": s["lead_id"],
        "email": s["email"],
        "company": s.get("company"),
        "variant": s.get("variant"),
        "subject": s.get("subject"),
        "message_id": None,
        "received_at": None,
        "recorded_at": datetime.now(timezone.utc).isoformat(),
        "kind": "reply",
        "note": note or "",
    })
    analytics.save_metrics(metrics)
    print(f"✅ Reply recorded for {s['contact_name']} <{s['email']}> ({s['company']})")


def cmd_replies():
    metrics = analytics.load_metrics()
    replies = metrics.get("replies", [])
    if not replies:
        print("No replies recorded yet. Use `record-reply <email|lead_id>` when one lands "
              "in the inbox, or set REPLY_TO to a Resend receiving domain to auto-detect.")
        return
    print(f"\n{'='*60}")
    print(f"  REPLIES ({len(replies)})")
    print(f"{'='*60}")
    for r in replies:
        note = f" — {r['note']}" if r.get("note") else ""
        source = "auto" if r.get("received_id") else "manual"
        print(f"  {r['recorded_at'][:16]}  {r['email']} ({r.get('company', '?')})  "
              f"[{r.get('variant', '?')} · {r.get('kind')} · {source}]{note}")
    print(f"{'='*60}\n")


# ── Experiment memory (the hourly feedback loop) ─────────────────────────────

def cmd_experiment(text: str = ""):
    """Append an entry to experiments/experiment_log.json: the hypothesis, the
    real current metrics (auto-attached), the send caps in force, and what the
    next batch should tweak. Called once per batch hour, before the next batch."""
    import analytics as _a
    log = load_sent_log()
    metrics = _a.load_metrics()
    t = _a.aggregate(log, metrics)
    cap, phase = daily_cap()
    entry = {
        "id": datetime.now(timezone.utc).strftime("EXP-%Y%m%d-%H%M"),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "provider": config.EMAIL_PROVIDER,
        "cap": {"daily_cap": cap, "phase": phase, "sent_today": sent_today(log)},
        "metrics": {
            "sent": t["sent"], "delivered": t["delivered"], "delivered_rate": round(t["delivered_rate"], 4),
            "opened": t["opened"], "open_rate": round(t["open_rate"], 4),
            "clicked": t["clicked"], "click_rate": round(t["click_rate"], 4),
            "replied": t["replied"], "reply_rate": round(t["reply_rate"], 4),
            "bounced": t["bounced"], "complained": t["complained"],
            "unverified": t["unknown"], "denied": t["denied"], "unresolved": t["unresolved"],
        },
        "by_variant": {v: {k: d[k] for k in ("sent", "open_rate", "reply_rate")}
                       for v, d in _a.by_variant(log, metrics).items()},
        "hypothesis": text,
    }
    config.EXPERIMENTS_DIR.mkdir(parents=True, exist_ok=True)
    path = config.EXPERIMENTS_DIR / "experiment_log.json"
    entries = []
    if path.exists():
        try:
            entries = json.loads(path.read_text())
        except Exception:
            entries = []
    entries.append(entry)
    path.write_text(json.dumps(entries, indent=2, default=str))
    print(f"✅ Experiment entry {entry['id']} appended to {path}")
    print(f"   reply {t['reply_rate']*100:.1f}% · open {t['open_rate']*100:.1f}% · "
          f"delivered {t['delivered_rate']*100:.1f}% · sent {t['sent']}")


# ── CLI ────────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="SpielOS Outbound Email Automation")
    parser.add_argument(
        "command",
        choices=["stats", "dry-run", "send", "metrics", "review", "record-reply", "replies", "experiment"],
        help="Action to perform",
    )
    parser.add_argument("--limit", type=int, default=None, help="Max emails to send")
    parser.add_argument("--lang", default=None, help="Language filter (en|fa or English|Persian)")
    parser.add_argument("--tier", default=None, help="Tier filter (A, B, C)")
    parser.add_argument("--retry-failed", action="store_true", help="send: retry the failed[] entries first")
    parser.add_argument("--force", action="store_true", help="metrics: bypass the schedule check")
    parser.add_argument("--quiet", action="store_true", help="metrics: one-line output (cron-friendly)")
    parser.add_argument("--json", action="store_true", help="review: JSON summary output")
    parser.add_argument("--note", default=None, help="record-reply: note about the reply")
    parser.add_argument("--text", default=None, help="experiment: hypothesis / tweak note for the log")
    parser.add_argument("identifier", nargs="?", default=None, help="record-reply: email or lead_id")
    args = parser.parse_args()

    lang = LANG_ALIASES.get((args.lang or "").strip().lower()) if args.lang else None
    if args.lang and not lang:
        print(f"ERROR: unknown --lang '{args.lang}' (en|fa)")
        sys.exit(1)

    if args.command == "stats":
        cmd_stats()
    elif args.command == "dry-run":
        cmd_dry_run(lang_filter=lang, tier_filter=args.tier, limit=args.limit or 5)
    elif args.command == "send":
        # English-first policy (2026-08-08): batches default to the English
        # list; Persian requires an explicit --lang fa until its copy quality
        # passes review.
        send_lang = lang or "English"
        if not args.limit and not args.retry_failed:
            print(f"Defaulting to a batch of {config.BATCH_SIZE} (--limit to override), language: {send_lang}")
        cmd_send(limit=args.limit, lang_filter=send_lang, tier_filter=args.tier, retry_failed=args.retry_failed)
    elif args.command == "metrics":
        cmd_metrics(force=args.force, quiet=args.quiet)
    elif args.command == "review":
        cmd_review(as_json=args.json)
    elif args.command == "record-reply":
        if not args.identifier:
            print("ERROR: record-reply needs an identifier (email or lead_id).")
            print("  Example: python3 outbound.py record-reply jane@acme.com")
            sys.exit(1)
        cmd_record_reply(args.identifier, note=args.note)
    elif args.command == "replies":
        cmd_replies()
    elif args.command == "experiment":
        cmd_experiment(text=args.text or "")


if __name__ == "__main__":
    main()
